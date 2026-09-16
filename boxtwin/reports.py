from io import BytesIO
from zoneinfo import ZoneInfo
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


NAVY = colors.HexColor("#071526")
BLUE = colors.HexColor("#1769FF")
CYAN = colors.HexColor("#18D4FF")
INK = colors.HexColor("#142238")
MUTED = colors.HexColor("#6F8095")
LINE = colors.HexColor("#DCE5EF")
PAPER = colors.HexColor("#F3F7FB")
GREEN = colors.HexColor("#147D64")
RED = colors.HexColor("#B12F3A")
FORTALEZA = ZoneInfo("America/Fortaleza")


def _local_datetime(value):
    if not value:
        return "-"
    try:
        parsed = __import__("datetime").datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=__import__("datetime").timezone.utc)
        return parsed.astimezone(FORTALEZA).strftime("%d/%m/%Y %H:%M")
    except (TypeError, ValueError):
        return str(value)


def _fmt(value, digits=1, suffix=""):
    if value is None:
        return "-"
    return f"{float(value):.{digits}f}{suffix}".replace(".", ",")


def build_operational_report(*, box_name, node_id, sensor_mode, dimensions, readings, anomalies, analytics=None):
    output = BytesIO()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("TitleBoxTwin", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=NAVY, spaceAfter=5)
    subtitle = ParagraphStyle("SubtitleBoxTwin", parent=styles["Normal"], fontSize=9, leading=13, textColor=MUTED)
    section = ParagraphStyle("SectionBoxTwin", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=12, leading=15, textColor=NAVY, spaceBefore=10, spaceAfter=8)
    body = ParagraphStyle("BodyBoxTwin", parent=styles["BodyText"], fontSize=8.5, leading=12, textColor=INK)
    tiny = ParagraphStyle("TinyBoxTwin", parent=body, fontSize=7, leading=9, textColor=MUTED)
    kpi_label = ParagraphStyle("KpiLabel", parent=body, alignment=TA_CENTER, fontSize=7.5, textColor=MUTED)
    kpi_value = ParagraphStyle("KpiValue", parent=body, alignment=TA_CENTER, fontName="Helvetica-Bold", fontSize=15, leading=18, textColor=NAVY)

    def footer(canvas, document):
        canvas.saveState()
        canvas.setStrokeColor(LINE)
        canvas.line(16 * mm, 13 * mm, 194 * mm, 13 * mm)
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(MUTED)
        canvas.drawString(16 * mm, 8 * mm, f"HydrogenI BoxTwin 3D - {node_id}")
        canvas.drawRightString(194 * mm, 8 * mm, f"Página {document.page}")
        canvas.restoreState()

    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=16 * mm,
        leftMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=18 * mm,
        title="Relatório operacional HydrogenI BoxTwin 3D",
        author="HydrogenI",
    )

    latest = readings[-1] if readings else None
    active_statuses = {"open", "acknowledged", "in_progress"}
    active_count = sum(item.get("status") in active_statuses for item in anomalies)
    avg_confidence = sum(item.get("confidence_percent", 0) for item in readings) / len(readings) if readings else None
    max_capacity = max((item.get("capacity_percent", 0) for item in readings), default=None)
    generated_at = __import__("datetime").datetime.now(FORTALEZA).strftime("%d/%m/%Y %H:%M")

    story = []
    header = Table([
        [Paragraph("H", ParagraphStyle("Mark", parent=title, alignment=TA_CENTER, textColor=CYAN, fontSize=24)),
         [Paragraph("RELATÓRIO OPERACIONAL", tiny), Paragraph("HydrogenI BoxTwin 3D", title), Paragraph(f"{box_name} | {node_id}", subtitle)]]
    ], colWidths=[18 * mm, 155 * mm])
    header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (0, 0), NAVY), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("BOX", (0, 0), (-1, -1), 0.7, LINE), ("LEFTPADDING", (1, 0), (1, 0), 7),
        ("TOPPADDING", (0, 0), (-1, -1), 8), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
    ]))
    story.extend([header, Spacer(1, 6 * mm)])

    status_text = "Sem leituras" if not latest else ("Alerta ativo" if latest.get("status") == "alert" else "Operação normal")
    analytics_kpis = (analytics or {}).get("kpis", {})
    kpis = [
        ("OCUPAÇÃO ATUAL", _fmt(latest.get("capacity_percent") if latest else None, 1, "%")),
        ("VOLUME EST. ATUAL", _fmt(analytics_kpis.get("current_volume_m3"), 4, " m³")),
        ("MASSA EST. ATUAL", _fmt(analytics_kpis.get("current_estimated_tons"), 3, " t")),
        ("INCIDENTES ATIVOS", str(active_count)),
    ]
    kpi_table = Table([[Table([[Paragraph(value, kpi_value)], [Paragraph(label, kpi_label)]]) for label, value in kpis]], colWidths=[43.25 * mm] * 4)
    kpi_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), PAPER), ("BOX", (0, 0), (-1, -1), .6, LINE),
        ("INNERGRID", (0, 0), (-1, -1), .6, colors.white), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
    ]))
    story.extend([kpi_table, Spacer(1, 5 * mm)])

    story.append(Paragraph("Resumo da operação", section))
    summary_data = [
        ["Gerado em", generated_at, "Modo do sensor", "Simulado" if sensor_mode == "mock" else "Físico"],
        ["Estado atual", status_text, "Leituras analisadas", str(len(readings))],
        ["Pico de ocupação", _fmt(max_capacity, 1, "%"), "Tempo médio de leitura", _fmt(analytics_kpis.get("average_reading_ms"), 0, " ms")],
        ["Desvios > +10%", str(analytics_kpis.get("above_expected_count", 0)), "Desvios < -10%", str(analytics_kpis.get("below_expected_count", 0))],
        ["Geometria de referência", f"{dimensions['length']} m x {dimensions['width']} m x {dimensions['height']} m", "Última leitura", _local_datetime(latest.get("created_at")) if latest else "-"],
    ]
    summary = Table(summary_data, colWidths=[31 * mm, 55.5 * mm, 31 * mm, 55.5 * mm])
    summary.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica"), ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("TEXTCOLOR", (0, 0), (-1, -1), INK), ("TEXTCOLOR", (0, 0), (0, -1), MUTED),
        ("TEXTCOLOR", (2, 0), (2, -1), MUTED), ("BACKGROUND", (0, 0), (-1, -1), colors.white),
        ("GRID", (0, 0), (-1, -1), .45, LINE), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))
    story.append(summary)

    story.append(Paragraph("Leituras recentes", section))
    reading_rows = [["Data e hora", "Box / sensor", "Material", "Volume (m³)", "t est.", "Desvio", "Estado"]]
    for item in list(reversed(readings))[:25]:
        reading_rows.append([
            _local_datetime(item.get("created_at")), f"{item.get('box_id') or node_id}\n{item.get('sensor_id') or node_id}", str(item.get("material_name") or item.get("material_type") or "-"), _fmt(item.get("volume_m3"), 4),
            _fmt(item.get("estimated_tons"), 3), _fmt(((item["volume_m3"] - item["expected_volume_m3"]) / item["expected_volume_m3"] * 100) if item.get("expected_volume_m3") else None, 1, "%"),
            "Alerta" if item.get("status") == "alert" else "Normal",
        ])
    if len(reading_rows) == 1:
        reading_rows.append(["Nenhuma leitura disponível", "-", "-", "-", "-", "-"])
    readings_table = Table(reading_rows, repeatRows=1, colWidths=[27 * mm, 31 * mm, 27 * mm, 24 * mm, 18 * mm, 19 * mm, 27 * mm])
    readings_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), .35, LINE),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PAPER]), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4.5), ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
    ]))
    story.append(readings_table)

    story.append(PageBreak())
    story.append(Paragraph("Incidentes e anomalias", section))
    status_labels = {"open": "Aberta", "acknowledged": "Ciente", "in_progress": "Em atendimento", "resolved": "Resolvida", "false_positive": "Falso positivo"}
    anomaly_rows = [["ID", "Data e hora", "Tipo", "Severidade", "Situação", "Descrição"]]
    for item in anomalies[:30]:
        anomaly_rows.append([
            f"#{item.get('id')}", _local_datetime(item.get("created_at")), str(item.get("anomaly_type") or "-"),
            "Crítica" if item.get("severity") == "danger" else "Atenção", status_labels.get(item.get("status"), item.get("status") or "-"),
            Paragraph(str(item.get("message") or "-"), tiny),
        ])
    if len(anomaly_rows) == 1:
        anomaly_rows.append(["-", "-", "-", "-", "-", "Nenhuma anomalia registrada"])
    anomaly_table = Table(anomaly_rows, repeatRows=1, colWidths=[11 * mm, 30 * mm, 23 * mm, 22 * mm, 29 * mm, 58 * mm])
    anomaly_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
        ("FONTSIZE", (0, 0), (-1, -1), 7), ("GRID", (0, 0), (-1, -1), .35, LINE),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PAPER]), ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(anomaly_table)
    story.extend([Spacer(1, 7 * mm), KeepTogether([
        Paragraph("Nota técnica", section),
        Paragraph(
            "O volume e a massa estimados representam a última leitura disponível de cada Box incluída no filtro; eles não somam capturas repetidas. A geometria apresentada é a configuração ativa do servidor no momento da exportação. No modo simulado, os valores demonstram o fluxo do MVP; a precisão industrial exige validação com sensor físico, instalação calibrada e volumes de referência conhecidos.",
            body,
        ),
    ])])

    document.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def build_operational_workbook(*, box_name, node_id, sensor_mode, dimensions, readings, anomalies, analytics=None):
    workbook = Workbook()
    navy, blue, white, pale = "071526", "1769FF", "FFFFFF", "F3F7FB"
    header_fill, pale_fill = PatternFill("solid", fgColor=navy), PatternFill("solid", fgColor=pale)
    header_font = Font(color=white, bold=True)

    def title_row(sheet, row, values):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row, column, value)
            cell.fill, cell.font = header_fill, header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def autosize(sheet):
        for column_cells in sheet.columns:
            letter = get_column_letter(column_cells[0].column)
            sheet.column_dimensions[letter].width = min(48, max(12, max(len(str(cell.value or "")) for cell in column_cells) + 2))

    summary_sheet = workbook.active
    summary_sheet.title = "Resumo"
    title_row(summary_sheet, 1, ["Projeto", "Box", "Node", "Sensor", "Gerado em"])
    summary_sheet.append(["HydrogenI BoxTwin 3D", box_name, node_id, "Simulado" if sensor_mode == "mock" else "Físico", datetime.now(FORTALEZA).replace(tzinfo=None)])
    summary_sheet["E2"].number_format = "dd/mm/yyyy hh:mm"
    summary_sheet.append([])
    title_row(summary_sheet, 4, ["Comprimento (m)", "Largura (m)", "Altura (m)", "Capacidade geométrica atual (m³)"])
    summary_sheet.append([dimensions["length"], dimensions["width"], dimensions["height"], dimensions["length"] * dimensions["width"] * dimensions["height"]])
    for cell in summary_sheet[5]: cell.number_format = "0.00000"
    summary_sheet.append([])
    title_row(summary_sheet, 7, ["Indicador", "Valor"])
    kpis = (analytics or {}).get("kpis", {})
    values = [
        ("Leituras no recorte", kpis.get("readings_count", len(readings))), ("Boxes no recorte", kpis.get("unique_boxes", 0)),
        ("Volume estimado atual (m³)", kpis.get("current_volume_m3", 0)), ("Massa estimada atual (t)", kpis.get("current_estimated_tons", 0)),
        ("Volume médio por leitura (m³)", kpis.get("average_volume_m3", 0)), ("Tempo médio de leitura (ms)", kpis.get("average_reading_ms")),
        ("Acima de +10% esperado", kpis.get("above_expected_count", 0)), ("Abaixo de -10% esperado", kpis.get("below_expected_count", 0)),
    ]
    for label, value in values: summary_sheet.append([label, value])
    for row in range(8, 16): summary_sheet.cell(row, 2).number_format = "0.00000"
    summary_sheet.freeze_panes = "A8"

    readings_sheet = workbook.create_sheet("Leituras")
    headers = ["ID", "UUID", "Data e hora", "Cenário", "Volume (m³)", "Ocupação (%)", "Confiança (%)", "Zonas válidas", "Estado", "Box", "Sensor", "Tipo de material", "Material", "Densidade (t/m³)", "Volume esperado (m³)", "Massa estimada (t)", "Tempo de leitura (ms)"]
    title_row(readings_sheet, 1, headers)
    for item in readings:
        created = item.get("created_at")
        try:
            parsed = datetime.fromisoformat(created.replace("Z", "+00:00")); created = parsed.astimezone(FORTALEZA).replace(tzinfo=None)
        except (AttributeError, ValueError):
            pass
        readings_sheet.append([item.get("id"), item.get("reading_uuid"), created, item.get("scenario"), item.get("volume_m3"), item.get("capacity_percent"), item.get("confidence_percent"), item.get("valid_zones"), item.get("status"), item.get("box_id"), item.get("sensor_id"), item.get("material_type"), item.get("material_name"), item.get("density_t_m3"), item.get("expected_volume_m3"), item.get("estimated_tons"), item.get("reading_duration_ms")])
    readings_sheet.freeze_panes, readings_sheet.auto_filter.ref = "A2", f"A1:Q{max(1, readings_sheet.max_row)}"
    for row in range(2, readings_sheet.max_row + 1):
        readings_sheet.cell(row, 3).number_format = "dd/mm/yyyy hh:mm"
        for col in (5, 14, 15, 16): readings_sheet.cell(row, col).number_format = "0.00000"
        for col in (6, 7): readings_sheet.cell(row, col).number_format = "0.0"

    anomalies_sheet = workbook.create_sheet("Tratativas")
    title_row(anomalies_sheet, 1, ["ID", "Data e hora", "Tipo", "Severidade", "Status", "Mensagem"])
    for item in anomalies:
        created = item.get("created_at")
        try:
            parsed = datetime.fromisoformat(created.replace("Z", "+00:00")); created = parsed.astimezone(FORTALEZA).replace(tzinfo=None)
        except (AttributeError, ValueError): pass
        anomalies_sheet.append([item.get("id"), created, item.get("anomaly_type"), item.get("severity"), item.get("status"), item.get("message")])
    anomalies_sheet.freeze_panes, anomalies_sheet.auto_filter.ref = "A2", f"A1:F{max(1, anomalies_sheet.max_row)}"
    for row in range(2, anomalies_sheet.max_row + 1): anomalies_sheet.cell(row, 2).number_format = "dd/mm/yyyy hh:mm"
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.row % 2 == 0: cell.fill = pale_fill
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        autosize(sheet)
    output = BytesIO(); workbook.save(output); return output.getvalue()
